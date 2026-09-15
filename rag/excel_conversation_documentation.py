"""Append completed customer-support conversations to a local Excel workbook."""

from __future__ import annotations

from datetime import datetime
from pathlib import Path
from typing import Any, Iterable, Optional

from openpyxl import Workbook, load_workbook

DEFAULT_EXCEL_DOCUMENTATION_PATH = (
    Path(__file__).resolve().parents[1] / "data/documentation/chatcom_conversations.xlsx"
)
WORKSHEET_NAME = "Conversations"
HEADERS = (
    "Conversation ID",
    "Timestamp",
    "Customer Message",
    "Agent Response",
    "Intent",
    "Sentiment",
    "Frustration Level",
    "Answerability",
    "Escalation Decision",
    "Escalation Reason",
    "Retrieved Source IDs",
    "Generation Mode",
    "Guardrail Result",
)


def append_conversation_to_excel(
    *,
    conversation_id: str,
    timestamp: str,
    customer_message: str,
    agent_response: str,
    intent: str,
    sentiment: str,
    frustration_level: str,
    answerability: str,
    escalation_decision: str,
    escalation_reason: Optional[str],
    retrieved_source_ids: Iterable[str],
    generation_mode: str,
    guardrail_result: str,
    destination: Path = DEFAULT_EXCEL_DOCUMENTATION_PATH,
) -> dict[str, Any]:
    """Append one completed conversation and return the write status."""

    destination = Path(destination)
    try:
        destination.parent.mkdir(parents=True, exist_ok=True)
        if destination.exists():
            workbook = load_workbook(destination)
            worksheet = workbook[WORKSHEET_NAME] if WORKSHEET_NAME in workbook.sheetnames else workbook.create_sheet(WORKSHEET_NAME)
        else:
            workbook = Workbook()
            worksheet = workbook.active
            worksheet.title = WORKSHEET_NAME

        _ensure_headers(worksheet)
        row = [
            conversation_id,
            timestamp,
            customer_message,
            agent_response,
            intent,
            sentiment,
            frustration_level,
            answerability,
            escalation_decision,
            escalation_reason or "",
            ", ".join(str(source_id) for source_id in retrieved_source_ids),
            generation_mode,
            guardrail_result,
        ]
        worksheet.append(row)
        workbook.save(destination)
        return {
            "success": True,
            "path": str(destination),
            "worksheet": WORKSHEET_NAME,
            "row": worksheet.max_row,
        }
    except Exception as exc:
        return {
            "success": False,
            "path": str(destination),
            "error": {"code": "excel_documentation_failed", "message": str(exc)},
        }


def _ensure_headers(worksheet: Any) -> None:
    existing_headers = tuple(worksheet.cell(row=1, column=index).value for index in range(1, len(HEADERS) + 1))
    if worksheet.max_row == 1 and all(value is None for value in existing_headers):
        for column, header in enumerate(HEADERS, start=1):
            worksheet.cell(row=1, column=column, value=header)
        return
    if existing_headers != HEADERS:
        raise ValueError(f"{WORKSHEET_NAME} worksheet must use the required headers")


def validate_excel_documentation(destination: Path = DEFAULT_EXCEL_DOCUMENTATION_PATH) -> dict[str, Any]:
    """Validate the workbook structure for tests and operational checks."""

    workbook = load_workbook(destination, read_only=True)
    worksheet = workbook[WORKSHEET_NAME]
    headers = tuple(cell.value for cell in worksheet[1][: len(HEADERS)])
    datetime.fromisoformat(str(worksheet.cell(row=2, column=2).value).replace("Z", "+00:00")) if worksheet.max_row > 1 else None
    return {"success": headers == HEADERS, "worksheet": WORKSHEET_NAME, "headers": headers}
