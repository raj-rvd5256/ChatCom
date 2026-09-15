"""Load the Excel business knowledge base into a persistent LangChain index."""

from __future__ import annotations

import os
from pathlib import Path
from typing import Any

from langchain_chroma import Chroma
from langchain_core.documents import Document
from langchain_huggingface import HuggingFaceEmbeddings
from openpyxl import load_workbook


PROJECT_ROOT = Path(__file__).resolve().parents[1]
DEFAULT_WORKBOOK = PROJECT_ROOT / "data" / "rag" / "ecommerce_knowledge_base.xlsx"
DEFAULT_INDEX = PROJECT_ROOT / "data" / "rag" / "chroma"
DEFAULT_EMBEDDING_MODEL = "sentence-transformers/all-MiniLM-L6-v2"
COLLECTION_NAME = "ecommerce_knowledge_base"
REQUIRED_COLUMNS = (
    "knowledge_id",
    "category",
    "topic",
    "question",
    "answer",
    "policy",
    "eligibility",
    "process",
    "timeline",
    "escalation_required",
    "source_reference",
)


def _load_documents(workbook_path: Path) -> list[Document]:
    workbook = load_workbook(workbook_path, read_only=True, data_only=True)
    if workbook.sheetnames != ["Knowledge_Base"]:
        raise ValueError("The workbook must contain only the Knowledge_Base worksheet.")

    sheet = workbook["Knowledge_Base"]
    rows = sheet.iter_rows(values_only=True)
    headers = tuple(next(rows))
    if headers != REQUIRED_COLUMNS:
        raise ValueError("The workbook headers do not match the required knowledge-base schema.")

    documents = []
    for row in rows:
        record = dict(zip(headers, row))
        if any(value is None or not str(value).strip() for value in record.values()):
            raise ValueError("The workbook contains a blank required field.")

        content = "\n".join(
            f"{column.replace('_', ' ').title()}: {record[column]}"
            for column in REQUIRED_COLUMNS
            if column not in {"knowledge_id", "source_reference"}
        )
        metadata: dict[str, Any] = {
            "knowledge_id": str(record["knowledge_id"]),
            "category": str(record["category"]),
            "topic": str(record["topic"]),
            "escalation_required": str(record["escalation_required"]),
            "source_reference": str(record["source_reference"]),
        }
        documents.append(Document(page_content=content, metadata=metadata))

    workbook.close()
    return documents


def _embeddings() -> HuggingFaceEmbeddings:
    model_name = os.getenv("RAG_EMBEDDING_MODEL", DEFAULT_EMBEDDING_MODEL)
    return HuggingFaceEmbeddings(model_name=model_name)


def build_vector_store(
    workbook_path: str | Path = DEFAULT_WORKBOOK,
    persist_directory: str | Path = DEFAULT_INDEX,
) -> Chroma:
    """Create or replace the persistent Chroma index from the Excel source."""
    workbook = Path(workbook_path)
    index = Path(persist_directory)
    documents = _load_documents(workbook)
    store = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=_embeddings(),
        persist_directory=str(index),
    )
    existing = store.get(include=[])
    if existing["ids"]:
        store.delete(ids=existing["ids"])
    store.add_documents(documents)
    return store


def get_retriever(
    persist_directory: str | Path = DEFAULT_INDEX,
    k: int = 4,
):
    """Return a similarity retriever over the persisted business knowledge."""
    store = Chroma(
        collection_name=COLLECTION_NAME,
        embedding_function=_embeddings(),
        persist_directory=str(persist_directory),
    )
    return store.as_retriever(search_kwargs={"k": k})


if __name__ == "__main__":
    vector_store = build_vector_store()
    print(f"Indexed {vector_store._collection.count()} knowledge records.")