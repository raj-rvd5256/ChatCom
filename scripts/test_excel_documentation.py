import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from rag.excel_conversation_documentation import HEADERS, WORKSHEET_NAME, append_conversation_to_excel
from rag.orchestrator import orchestrate_customer_support


class FakeRetriever:
    def invoke(self, question):
        return [
            SimpleNamespace(
                page_content=(
                    "Category: Shipping | Topic: Shipping method | "
                    "Question: What shipping method is available? | "
                    "Answer: Standard shipping is available."
                ),
                metadata={
                    "knowledge_id": "NK-001",
                    "category": "Shipping",
                    "topic": "Shipping method",
                    "source_reference": "Knowledge_Base!NK-001",
                    "escalation_required": "No",
                },
            )
        ]


def conversation_kwargs():
    return {
        "conversation_id": "conversation-001",
        "timestamp": "2026-09-15T12:00:00Z",
        "customer_message": "What shipping method is available?",
        "agent_response": "Standard shipping is available.",
        "intent": "shipping",
        "sentiment": "neutral",
        "frustration_level": "low",
        "answerability": "directly_answerable",
        "escalation_decision": "not_required",
        "escalation_reason": None,
        "retrieved_source_ids": ["NK-001"],
        "generation_mode": "deterministic_retrieved_knowledge",
        "guardrail_result": "passed",
    }


class ExcelConversationDocumentationTests(unittest.TestCase):
    def test_creates_workbook_sheet_headers_and_first_row(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chatcom_conversations.xlsx"
            result = append_conversation_to_excel(destination=path, **conversation_kwargs())
            workbook = load_workbook(path)
            worksheet = workbook[WORKSHEET_NAME]

            self.assertTrue(result["success"])
            self.assertEqual(tuple(cell.value for cell in worksheet[1]), HEADERS)
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(worksheet.cell(row=2, column=1).value, "conversation-001")

    def test_preserves_existing_rows_and_appends_second_conversation(self):
        with tempfile.TemporaryDirectory() as directory:
            path = Path(directory) / "chatcom_conversations.xlsx"
            append_conversation_to_excel(destination=path, **conversation_kwargs())
            second = conversation_kwargs()
            second["conversation_id"] = "conversation-002"
            second["customer_message"] = "How long does shipping take?"
            append_conversation_to_excel(destination=path, **second)
            worksheet = load_workbook(path)[WORKSHEET_NAME]

            self.assertEqual(worksheet.max_row, 3)
            self.assertEqual(worksheet.cell(row=2, column=1).value, "conversation-001")
            self.assertEqual(worksheet.cell(row=3, column=1).value, "conversation-002")

    def test_orchestrator_documents_conversation_in_excel(self):
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "chatcom_conversations.xlsx"
            result = orchestrate_customer_support(
                "What shipping method is available?",
                retriever=FakeRetriever(),
                documentation_destination=Path(directory) / "evidence.jsonl",
                excel_documentation_destination=excel_path,
            )
            worksheet = load_workbook(excel_path)[WORKSHEET_NAME]

            self.assertTrue(result["excel_documentation_result"]["success"])
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(worksheet.cell(row=2, column=3).value, "What shipping method is available?")

    def test_excel_failure_does_not_break_customer_response(self):
        with tempfile.TemporaryDirectory() as directory:
            with patch(
                "rag.orchestrator.append_conversation_to_excel",
                side_effect=OSError("workbook unavailable"),
            ):
                result = orchestrate_customer_support(
                    "What shipping method is available?",
                    retriever=FakeRetriever(),
                    documentation_destination=Path(directory) / "evidence.jsonl",
                    excel_documentation_destination=Path(directory) / "chatcom_conversations.xlsx",
                )

            self.assertIn("Standard shipping is available", result["response"])
            self.assertFalse(result["excel_documentation_result"]["success"])
            self.assertTrue(result["documentation_result"]["success"])


if __name__ == "__main__":
    unittest.main()
