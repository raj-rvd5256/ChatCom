import contextlib
import io
import tempfile
import unittest
from pathlib import Path
from types import SimpleNamespace
from unittest.mock import patch

from openpyxl import load_workbook

from rag.excel_conversation_documentation import WORKSHEET_NAME
from rag.orchestrator import orchestrate_customer_support
from scripts import run_customer_support


class CustomerSupportAdapterTests(unittest.TestCase):
    def fake_result(self, response="A grounded answer."):
        return {
            "response": response,
            "intent": "shipping",
            "sentiment": "neutral",
            "frustration_level": "low",
            "answerable": True,
            "answerability_state": "directly_answerable",
            "escalation_required": False,
            "escalation_reason": None,
            "retrieved_knowledge_ids": ["NK-001"],
            "generation_mode": "deterministic_retrieved_knowledge",
            "guardrail_result": "passed",
            "excel_documentation_result": {"success": True},
            "documentation_result": {"success": True},
        }

    def test_adapter_invokes_orchestrator_and_hides_metadata(self):
        output = io.StringIO()
        with patch.object(run_customer_support, "orchestrate_customer_support", return_value=self.fake_result()) as orchestrator:
            with contextlib.redirect_stdout(output):
                exit_code = run_customer_support.main(["What", "is", "shipping?"])

        orchestrator.assert_called_once_with("What is shipping?")
        self.assertEqual(exit_code, 0)
        self.assertEqual(output.getvalue().strip(), "A grounded answer.")
        self.assertNotIn("Intent:", output.getvalue())
        self.assertNotIn("NK-001", output.getvalue())

    def test_customer_question_creates_excel_row(self):
        with tempfile.TemporaryDirectory() as directory:
            excel_path = Path(directory) / "chatcom_conversations.xlsx"
            jsonl_path = Path(directory) / "evidence.jsonl"

            def invoke(question):
                return orchestrate_customer_support(
                    question,
                    retriever=FakeRetriever(),
                    documentation_destination=jsonl_path,
                    excel_documentation_destination=excel_path,
                )

            with patch.object(run_customer_support, "orchestrate_customer_support", side_effect=invoke) as orchestrator:
                exit_code = run_customer_support.main(["What", "is", "shipping?"])

            worksheet = load_workbook(excel_path, read_only=True)[WORKSHEET_NAME]
            self.assertEqual(exit_code, 0)
            self.assertEqual(orchestrator.call_count, 1)
            self.assertEqual(worksheet.max_row, 2)
            self.assertEqual(worksheet.cell(row=2, column=3).value, "What is shipping?")

    def test_evaluation_mode_exposes_required_metadata(self):
        output = io.StringIO()
        with patch.object(run_customer_support, "orchestrate_customer_support", return_value=self.fake_result()):
            with contextlib.redirect_stdout(output):
                exit_code = run_customer_support.main(["What is shipping?", "--evaluation"])

        rendered = output.getvalue()
        self.assertEqual(exit_code, 0)
        for label in (
            "Customer-facing response:",
            "Intent:",
            "Sentiment:",
            "Frustration level:",
            "Answerability:",
            "Escalation decision:",
            "Escalation reason:",
            "Retrieved source IDs:",
            "Documentation status:",
            "Generation mode:",
            "Guardrail result:",
        ):
            self.assertIn(label, rendered)

    def test_excel_failure_does_not_break_adapter_response(self):
        result = self.fake_result("Still a grounded answer.")
        result["excel_documentation_result"] = {
            "success": False,
            "error": {"code": "excel_documentation_failed"},
        }
        output = io.StringIO()
        with patch.object(run_customer_support, "orchestrate_customer_support", return_value=result):
            with contextlib.redirect_stdout(output):
                exit_code = run_customer_support.main(["What is shipping?"])

        self.assertEqual(exit_code, 0)
        self.assertEqual(output.getvalue().strip(), "Still a grounded answer.")
        self.assertNotIn("excel_documentation_failed", output.getvalue())


class FakeRetriever:
    def invoke(self, question):
        return [
            SimpleNamespace(
                page_content=(
                    "Category: Shipping | Topic: Shipping method | "
                    "Question: What is shipping? | Answer: Standard shipping is available."
                ),
                metadata={
                    "knowledge_id": "NK-001",
                    "category": "Shipping",
                    "topic": "Shipping method",
                    "source_reference": "Knowledge_Base!NK-001",
                    "escalation_required": "No",
                },
            )
        ]


if __name__ == "__main__":
    unittest.main()
